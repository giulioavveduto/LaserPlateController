from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from experiment.run_report import build_run_rows
from plates.plate_geometry import PlateGeometry

if TYPE_CHECKING:
    from experiment.experiment_runner import ExperimentRunner


RESULT_COLUMNS = [
    ("sequence", "Sequence"),
    ("well", "Well"),
    ("outcome", "Outcome"),
    ("planned_duration_s", "Planned duration (s)"),
    ("executed_duration_s", "Executed duration (s)"),
    ("completion_percent", "Completion (%)"),
    ("requested_mode", "Requested setpoint mode"),
    ("requested_value", "Requested setpoint value"),
    ("requested_unit", "Requested unit"),
    ("programmed_current_percent", "Programmed current (%)"),
    ("estimated_power_w", "Estimated power (W)"),
    (
        "estimated_irradiance_w_cm2",
        "Estimated irradiance (W/cm²)",
    ),
    ("executed_energy_j", "Executed energy (J)"),
    (
        "executed_fluence_j_cm2",
        "Executed fluence (J/cm²)",
    ),
    ("calibration_id", "Calibration ID"),
    ("execution_mode", "Execution mode"),
    ("terminal_state", "Terminal state"),
    (
        "final_laser_off_confirmed",
        "Final laser OFF confirmed",
    ),
]

SETPOINT_UNITS = {
    "current_percent": "%",
    "power_w": "W",
    "irradiance_w_cm2": "W/cm²",
}

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)
HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def _load_snapshot(
    run_directory: Path,
) -> dict[str, object]:
    snapshot_path = run_directory / "protocol.lpp"

    data = json.loads(snapshot_path.read_text(encoding="utf-8"))

    if not isinstance(data, dict):
        raise ValueError("The run protocol snapshot is invalid.")

    if not isinstance(data.get("run"), dict):
        raise ValueError("The run metadata is missing from protocol.lpp.")

    return data


def build_excel_rows(
    runner: ExperimentRunner,
    snapshot: dict[str, object],
) -> list[dict[str, object]]:
    run_data = snapshot["run"]
    resolved_by_well = run_data.get(
        "resolved_laser_setpoints_by_well",
        {},
    )

    if not isinstance(resolved_by_well, dict):
        resolved_by_well = {}

    rows: list[dict[str, object]] = []

    for base_row in build_run_rows(runner):
        well = str(base_row["well"])
        resolved = resolved_by_well.get(well, {})

        if not isinstance(resolved, dict):
            resolved = {}

        executed_s = float(base_row["executed_duration_s"])

        estimated_power = resolved.get("estimated_power_w")
        estimated_irradiance = resolved.get("estimated_irradiance_w_cm2")

        if estimated_power is not None:
            estimated_power = float(estimated_power)

        if estimated_irradiance is not None:
            estimated_irradiance = float(estimated_irradiance)

        requested_mode = str(resolved.get("requested_mode", ""))

        rows.append(
            {
                "sequence": int(base_row["sequence"]),
                "well": well,
                "outcome": base_row["outcome"],
                "planned_duration_s": float(base_row["planned_duration_s"]),
                "executed_duration_s": executed_s,
                "completion_percent": float(base_row["completion_percent"]),
                "requested_mode": requested_mode,
                "requested_value": resolved.get("requested_value"),
                "requested_unit": SETPOINT_UNITS.get(
                    requested_mode,
                    "",
                ),
                "programmed_current_percent": (base_row["programmed_current_percent"]),
                "estimated_power_w": estimated_power,
                "estimated_irradiance_w_cm2": (estimated_irradiance),
                "executed_energy_j": (
                    None if estimated_power is None else estimated_power * executed_s
                ),
                "executed_fluence_j_cm2": (
                    None
                    if estimated_irradiance is None
                    else estimated_irradiance * executed_s
                ),
                "calibration_id": resolved.get("calibration_id") or "",
                "execution_mode": base_row["execution_mode"],
                "terminal_state": base_row["terminal_state"],
                "final_laser_off_confirmed": (base_row["final_laser_off_confirmed"]),
            }
        )

    return rows


def _style_header(
    worksheet,
    row_number: int,
) -> None:
    for cell in worksheet[row_number]:
        if cell.value is None:
            continue

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _fit_columns(
    worksheet,
    maximum_width: int = 38,
) -> None:
    for column_cells in worksheet.columns:
        width = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = (
            min(max(width + 2, 10), maximum_width)
        )


def write_excel_report(
    run_directory: str | Path,
    runner: ExperimentRunner,
) -> Path:
    directory = Path(run_directory)
    snapshot = _load_snapshot(directory)
    run_data = snapshot["run"]

    plate = PlateGeometry(str(snapshot.get("plate_type", "")))
    result_rows = build_excel_rows(
        runner,
        snapshot,
    )

    workbook = Workbook()

    results_sheet = workbook.active
    results_sheet.title = "Well results"
    results_sheet.append([label for _, label in RESULT_COLUMNS])

    for row in result_rows:
        results_sheet.append([row[field_name] for field_name, _ in RESULT_COLUMNS])

    _style_header(results_sheet, 1)
    results_sheet.freeze_panes = "A2"
    results_sheet.auto_filter.ref = results_sheet.dimensions

    three_decimal_headers = {
        "Planned duration (s)",
        "Executed duration (s)",
        "Requested setpoint value",
        "Estimated power (W)",
        "Estimated irradiance (W/cm²)",
        "Executed energy (J)",
        "Executed fluence (J/cm²)",
    }

    for column_index, (_, label) in enumerate(
        RESULT_COLUMNS,
        start=1,
    ):
        number_format = (
            "0.0"
            if label == "Completion (%)"
            else "0.000" if label in three_decimal_headers else None
        )

        if number_format is not None:
            for cell in results_sheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
            ):
                for item in cell:
                    item.number_format = number_format

    _fit_columns(results_sheet)

    metadata_sheet = workbook.create_sheet("Run metadata")

    a1_mm = run_data.get("a1_mm")
    ordering_start = run_data.get("ordering_start_position_mm")

    metadata = [
        ("Report format version", 1),
        ("Run ID", run_data.get("id", "")),
        ("Experiment name", snapshot.get("name", "")),
        ("Started at", run_data.get("created_at", "")),
        (
            "Report created at",
            datetime.now().astimezone().isoformat(timespec="seconds"),
        ),
        ("Terminal state", runner.state.name),
        ("Plate type", plate.name),
        ("Well area (cm²)", plate.well_area_cm2),
        ("Execution mode", run_data.get("mode", "")),
        ("Stage mode", run_data.get("stage_mode", "")),
        ("Laser mode", run_data.get("laser_mode", "")),
        (
            "Irradiation order",
            run_data.get(
                "irradiation_order",
                snapshot.get(
                    "irradiation_order",
                    "row",
                ),
            ),
        ),
        (
            "Resolved well sequence",
            ", ".join(
                run_data.get(
                    "resolved_well_sequence",
                    [],
                )
            ),
        ),
        (
            "A1 X (mm)",
            (a1_mm[0] if isinstance(a1_mm, list) and len(a1_mm) == 2 else None),
        ),
        (
            "A1 Y (mm)",
            (a1_mm[1] if isinstance(a1_mm, list) and len(a1_mm) == 2 else None),
        ),
        (
            "Ordering start X (mm)",
            (
                ordering_start[0]
                if isinstance(ordering_start, list) and len(ordering_start) == 2
                else None
            ),
        ),
        (
            "Ordering start Y (mm)",
            (
                ordering_start[1]
                if isinstance(ordering_start, list) and len(ordering_start) == 2
                else None
            ),
        ),
        (
            "Final laser OFF confirmed",
            (
                "Not applicable"
                if runner.stage_only
                else ("Yes" if runner.last_laser_off_confirmed else "No")
            ),
        ),
    ]

    metadata_sheet.append(["Run property", "Value"])

    for key, value in metadata:
        metadata_sheet.append([key, value])

    _style_header(metadata_sheet, 1)

    calibration_header_row = metadata_sheet.max_row + 2
    metadata_sheet.cell(
        calibration_header_row,
        1,
        "Calibration ID",
    )
    metadata_sheet.cell(
        calibration_header_row,
        2,
        "Calibration date (UTC)",
    )
    metadata_sheet.cell(
        calibration_header_row,
        3,
        "Current (%)",
    )
    metadata_sheet.cell(
        calibration_header_row,
        4,
        "Measured power (W)",
    )
    _style_header(
        metadata_sheet,
        calibration_header_row,
    )

    calibrations = run_data.get(
        "laser_calibrations",
        {},
    )

    if isinstance(calibrations, dict):
        for calibration_id, calibration in calibrations.items():
            if not isinstance(calibration, dict):
                continue

            points = calibration.get("points", [])

            if not isinstance(points, list):
                continue

            for point in points:
                if not isinstance(point, dict):
                    continue

                metadata_sheet.append(
                    [
                        calibration_id,
                        calibration.get(
                            "created_at_utc",
                            "",
                        ),
                        point.get("current_percent"),
                        point.get("power_w"),
                    ]
                )

    for row in metadata_sheet.iter_rows(
        min_row=2,
        max_row=metadata_sheet.max_row,
    ):
        if row[0].value in {
            "Well area (cm²)",
            "A1 X (mm)",
            "A1 Y (mm)",
            "Ordering start X (mm)",
            "Ordering start Y (mm)",
        }:
            row[1].number_format = "0.000"

    for row in metadata_sheet.iter_rows(
        min_row=calibration_header_row + 1,
        min_col=4,
        max_col=4,
    ):
        row[0].number_format = "0.000"

    metadata_sheet.freeze_panes = "A2"
    _fit_columns(metadata_sheet, maximum_width=60)

    report_path = directory / "run_report.xlsx"
    workbook.save(report_path)

    return report_path
