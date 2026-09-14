from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from experiment.excel_report import (
    build_excel_rows,
    write_excel_report,
)
from experiment.experiment_runner import ExperimentState


class ExcelReportTests(unittest.TestCase):
    def make_runner(self) -> SimpleNamespace:
        return SimpleNamespace(
            wells=["A1", "C3"],
            exposure_times_s=[5.0, 10.0],
            executed_times_s={
                "A1": 5.0,
                "C3": 3.5,
            },
            completed_wells=["A1"],
            stage_only=False,
            current_percents=[30, 55],
            last_laser_off_confirmed=True,
            state=ExperimentState.STOPPED,
        )

    def make_snapshot(self) -> dict[str, object]:
        calibration_id = "calibration-test-001"

        return {
            "format_version": 2,
            "name": "Excel report test",
            "plate_type": "96-well plate",
            "selected_wells": ["A1", "C3"],
            "irradiation_order": "optimized",
            "common_exposure_time_s": 5.0,
            "default_laser_setpoint": None,
            "well_treatments": {},
            "run": {
                "id": "run-test-001",
                "created_at": ("2026-09-14T10:00:00+02:00"),
                "mode": "automatic_irradiation",
                "stage_mode": "Simulator",
                "laser_mode": "Simulator",
                "a1_mm": [12.5, 18.0],
                "irradiation_order": "optimized",
                "resolved_well_sequence": [
                    "A1",
                    "C3",
                ],
                "ordering_start_position_mm": [
                    0.0,
                    0.0,
                ],
                "resolved_laser_setpoints_by_well": {
                    "A1": {
                        "requested_mode": "power_w",
                        "requested_value": 0.5,
                        "calibration_id": calibration_id,
                        "current_percent": 30,
                        "estimated_power_w": 0.49,
                        "estimated_irradiance_w_cm2": 0.25,
                    },
                    "C3": {
                        "requested_mode": ("irradiance_w_cm2"),
                        "requested_value": 0.4,
                        "calibration_id": calibration_id,
                        "current_percent": 55,
                        "estimated_power_w": 0.8,
                        "estimated_irradiance_w_cm2": 0.39,
                    },
                },
                "laser_calibrations": {
                    calibration_id: {
                        "calibration_id": calibration_id,
                        "created_at_utc": ("2026-09-14T08:00:00Z"),
                        "points": [
                            {
                                "current_percent": 10,
                                "power_w": 0.1,
                            },
                            {
                                "current_percent": 30,
                                "power_w": 0.49,
                            },
                            {
                                "current_percent": 55,
                                "power_w": 0.8,
                            },
                        ],
                    }
                },
            },
        }

    def test_rows_include_executed_dose(self) -> None:
        rows = build_excel_rows(
            self.make_runner(),
            self.make_snapshot(),
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[0]["requested_unit"],
            "W",
        )
        self.assertAlmostEqual(
            rows[0]["executed_energy_j"],
            2.45,
        )
        self.assertAlmostEqual(
            rows[0]["executed_fluence_j_cm2"],
            1.25,
        )
        self.assertEqual(
            rows[1]["requested_unit"],
            "W/cm²",
        )
        self.assertAlmostEqual(
            rows[1]["executed_energy_j"],
            2.8,
        )
        self.assertAlmostEqual(
            rows[1]["executed_fluence_j_cm2"],
            1.365,
        )

    def test_workbook_contains_results_and_metadata(
        self,
    ) -> None:
        runner = self.make_runner()
        snapshot = self.make_snapshot()

        with tempfile.TemporaryDirectory() as directory:
            run_directory = Path(directory)
            snapshot_path = run_directory / "protocol.lpp"
            snapshot_path.write_text(
                json.dumps(snapshot),
                encoding="utf-8",
            )

            report_path = write_excel_report(
                run_directory,
                runner,
            )

            self.assertEqual(
                report_path,
                run_directory / "run_report.xlsx",
            )
            self.assertTrue(report_path.exists())

            workbook = load_workbook(
                report_path,
                data_only=True,
            )

            self.assertEqual(
                workbook.sheetnames,
                [
                    "Well results",
                    "Run metadata",
                ],
            )

            results = workbook["Well results"]
            headers = {cell.value: cell.column for cell in results[1]}

            self.assertEqual(
                results.cell(
                    2,
                    headers["Well"],
                ).value,
                "A1",
            )
            self.assertAlmostEqual(
                results.cell(
                    2,
                    headers["Executed energy (J)"],
                ).value,
                2.45,
            )

            metadata_sheet = workbook["Run metadata"]
            metadata = {
                row[0].value: row[1].value
                for row in metadata_sheet.iter_rows(
                    min_row=2,
                    max_col=2,
                )
                if row[0].value is not None
            }

            self.assertEqual(
                metadata["Plate type"],
                "96-well plate",
            )
            self.assertEqual(
                metadata["Irradiation order"],
                "optimized",
            )
            self.assertEqual(
                metadata["Laser mode"],
                "Simulator",
            )

            calibration_rows = [
                tuple(cell.value for cell in row[:4])
                for row in metadata_sheet.iter_rows()
                if row[0].value == "calibration-test-001"
            ]

            self.assertEqual(
                len(calibration_rows),
                3,
            )
            self.assertEqual(
                calibration_rows[0][2:],
                (10, 0.1),
            )


if __name__ == "__main__":
    unittest.main()
